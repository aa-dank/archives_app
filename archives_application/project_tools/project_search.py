"""Read-only project metadata search and spreadsheet-export preparation."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO
from itertools import islice
import re

import flask
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import and_, case, exists, false, func, literal, or_

from archives_application import db, utils
from archives_application.models import CAANModel, ContractModel, ProjectModel


HTML_RESULT_LIMIT = 300
SEARCH_PARAMETERS = frozenset({"query", "status", "drawings", "has_archive_location"})
FILTER_VALUES = {
    "status": frozenset({"any", "open", "closed", "unknown"}),
    "drawings": frozenset({"any", "yes", "no", "yes_or_unknown"}),
    "has_archive_location": frozenset({"any", "yes", "no"}),
}
PROJECT_FIELDS = (
    ProjectModel.number,
    ProjectModel.name,
    ProjectModel.project_manager_name,
    ProjectModel.inspector_name,
)
CONTRACT_FIELDS = (
    ContractModel.contract_number,
    ContractModel.contractor_org_name,
    ContractModel.executive_design_org_name,
    ContractModel.scope_description,
)


class ProjectSearchValidationError(ValueError):
    """Raised when a project-search request does not meet its public contract."""


@dataclass(frozen=True)
class ProjectSearchState:
    """Validated, presentation-ready state for a project-search request."""

    query: str
    terms: tuple[str, ...]
    status: str
    drawings: str
    has_archive_location: str

    @property
    def has_active_filter(self) -> bool:
        return any(
            value != "any"
            for value in (self.status, self.drawings, self.has_archive_location)
        )

    @property
    def is_active(self) -> bool:
        return bool(self.query) or self.has_active_filter

    def export_parameters(self) -> dict[str, str]:
        """Return the minimal parameter set that preserves this active search."""
        values = {}
        if self.query:
            values["query"] = self.query
        for name in ("status", "drawings", "has_archive_location"):
            value = getattr(self, name)
            if value != "any":
                values[name] = value
        return values

    def active_filter_labels(self) -> list[tuple[str, str]]:
        labels = {
            "status": "Status",
            "drawings": "Drawings",
            "has_archive_location": "File server location",
        }
        display_values = {
            "yes_or_unknown": "Yes or Unknown",
            "has_archive_location:yes": "Known",
            "has_archive_location:no": "Unknown",
        }
        values = []
        for name in ("status", "drawings", "has_archive_location"):
            value = getattr(self, name)
            if value != "any":
                display_key = f"{name}:{value}"
                values.append((
                    labels[name],
                    display_values.get(display_key, display_values.get(value, value.title())),
                ))
        return values


@dataclass(frozen=True)
class ProjectSearchResult:
    """One ranked project row plus contract aggregate data."""

    project: ProjectModel
    contract_count: int
    recorded_cost_count: int
    original_cost_total: Decimal | None
    ranking_score: int
    matched_project: bool
    matched_contract: bool
    matched_caan: bool

    @property
    def ranking_band(self) -> str:
        return {
            1: "Exact project number",
            2: "Exact linked CAAN",
            3: "Project-number prefix",
            4: "Exact project name",
            5: "Project metadata",
            6: "One linked contract",
            7: "Distributed metadata match",
            8: "Filter-only result",
        }[self.ranking_score]

    @property
    def matched_in(self) -> str:
        labels = []
        if self.matched_project:
            labels.append("Project")
        if self.matched_contract:
            labels.append("Contract")
        if self.matched_caan:
            labels.append("CAAN")
        return ", ".join(labels) if labels else "—"

    @property
    def initial_contract_value(self) -> str:
        if self.contract_count == 0:
            return "No contract data"
        if self.recorded_cost_count == 0:
            return "Not recorded"
        total = f"${self.original_cost_total:,.2f}"
        return total if self.recorded_cost_count == self.contract_count else f"{total} (Partial)"


def parse_request(query_args) -> ProjectSearchState:
    """Validate public query parameters without treating a query as a selector."""
    supplied_keys = set(query_args.keys())
    unknown_keys = supplied_keys - SEARCH_PARAMETERS
    if unknown_keys:
        raise ProjectSearchValidationError(
            "Unknown query parameter(s): " + ", ".join(sorted(unknown_keys))
        )
    for name in supplied_keys:
        if len(query_args.getlist(name)) > 1:
            raise ProjectSearchValidationError(f"{name} must be supplied only once.")

    raw_query = query_args.get("query", "")
    query = raw_query.strip()
    if len(query) > 200:
        raise ProjectSearchValidationError("query must be 200 characters or fewer.")

    filters = {}
    for name, allowed_values in FILTER_VALUES.items():
        raw_value = query_args.get(name, "any")
        value = raw_value.strip().lower()
        if value not in allowed_values:
            raise ProjectSearchValidationError(
                f"{name} must be one of: " + ", ".join(sorted(allowed_values)) + "."
            )
        filters[name] = value

    return ProjectSearchState(
        query=query,
        terms=tuple(term.lower() for term in query.split() if term),
        **filters,
    )


def _like_pattern(value: str, suffix: bool = True) -> str:
    """Escape a literal ILIKE value while retaining PostgreSQL's backslash escape."""
    escaped = _escape_like_value(value)
    return f"%{escaped}%" if suffix else f"{escaped}%"


def _escape_like_value(value: str) -> str:
    return value.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_")


def _field_match(fields, term: str):
    return or_(*(field.ilike(_like_pattern(term), escape="\\") for field in fields))


def _funding_number_token_match(term: str):
    """Match one whole funding identifier, including values stored across lines."""
    normalized_funding = func.lower(func.coalesce(ContractModel.funding_number, ""))
    for whitespace_character in ("\n", "\r", "\t"):
        normalized_funding = func.replace(
            normalized_funding, whitespace_character, " "
        )
    padded_funding = literal(" ") + normalized_funding + literal(" ")
    return padded_funding.like(f"% {_escape_like_value(term)} %", escape="\\")


def _contract_exists(predicate):
    return exists().where(and_(ContractModel.project_id == ProjectModel.id, predicate))


def _contract_summary_subquery():
    return (
        db.session.query(
            ContractModel.project_id.label("project_id"),
            func.count(ContractModel.id).label("contract_count"),
            func.count(ContractModel.original_contract_cost).label("recorded_cost_count"),
            func.sum(ContractModel.original_contract_cost).label("original_cost_total"),
        )
        .filter(ContractModel.project_id.isnot(None))
        .group_by(ContractModel.project_id)
        .subquery()
    )


def _ranked_query(state: ProjectSearchState):
    """Return the one-project-per-row ranked query shared by HTML and export."""
    summary = _contract_summary_subquery()
    normalized_number = func.lower(func.btrim(ProjectModel.number))
    normalized_name = func.lower(func.btrim(ProjectModel.name))
    normalized_caan = func.lower(func.btrim(CAANModel.caan))

    if state.terms:
        local_terms = [_field_match(PROJECT_FIELDS, term) for term in state.terms]
        contract_terms = [
            or_(
                _field_match(CONTRACT_FIELDS, term),
                _funding_number_token_match(term),
            )
            for term in state.terms
        ]
        contract_term_exists = [_contract_exists(predicate) for predicate in contract_terms]
        caan_term_exists = [
            ProjectModel.caans.any(normalized_caan == term)
            for term in state.terms
        ]
        all_local = and_(*local_terms)
        one_contract_matches_all = _contract_exists(and_(*contract_terms))
        matched_project = or_(*local_terms)
        matched_contract = or_(*contract_term_exists)
        matched_caan = or_(*caan_term_exists)
        # CAAN discovery is identifier-only. Each term can match a direct CAAN
        # code exactly, and terms may be supplied by different linked records.
        exact_linked_caan = ProjectModel.caans.any(
            normalized_caan == state.query.lower()
        )
        # A term may be supplied by project metadata or any direct contract.
        # Each EXISTS is intentionally independent so terms may be distributed.
        match_predicates = [
            or_(local_match, contract_match, caan_match)
            for local_match, contract_match, caan_match in zip(
                local_terms, contract_term_exists, caan_term_exists
            )
        ]
        normalized_query = state.query.lower()
        exact_number = normalized_number == normalized_query
        number_prefix = normalized_number.ilike(_like_pattern(normalized_query, suffix=False), escape="\\")
        exact_name = normalized_name == normalized_query
        ranking_score = case(
            (exact_number, 1),
            (exact_linked_caan, 2),
            (number_prefix, 3),
            (exact_name, 4),
            (all_local, 5),
            (one_contract_matches_all, 6),
            else_=7,
        )
    else:
        match_predicates = []
        matched_project = false()
        matched_contract = false()
        matched_caan = false()
        exact_linked_caan = false()
        ranking_score = literal(8)

    query = (
        db.session.query(
            ProjectModel,
            func.coalesce(summary.c.contract_count, 0).label("contract_count"),
            func.coalesce(summary.c.recorded_cost_count, 0).label("recorded_cost_count"),
            summary.c.original_cost_total.label("original_cost_total"),
            ranking_score.label("ranking_score"),
            matched_project.label("matched_project"),
            matched_contract.label("matched_contract"),
            matched_caan.label("matched_caan"),
        )
        .outerjoin(summary, summary.c.project_id == ProjectModel.id)
    )
    if match_predicates:
        query = query.filter(or_(and_(*match_predicates), exact_linked_caan))

    if state.status == "open":
        query = query.filter(ProjectModel.closed.is_(False))
    elif state.status == "closed":
        query = query.filter(ProjectModel.closed.is_(True))
    elif state.status == "unknown":
        query = query.filter(ProjectModel.closed.is_(None))

    if state.drawings == "yes":
        query = query.filter(ProjectModel.drawings.is_(True))
    elif state.drawings == "no":
        query = query.filter(ProjectModel.drawings.is_(False))
    elif state.drawings == "yes_or_unknown":
        query = query.filter(or_(ProjectModel.drawings.is_(True), ProjectModel.drawings.is_(None)))

    trimmed_archive_root = func.btrim(ProjectModel.file_server_location, " \t\r\n")
    root_is_recorded = and_(
        ProjectModel.file_server_location.isnot(None),
        trimmed_archive_root != "",
    )
    if state.has_archive_location == "yes":
        query = query.filter(root_is_recorded)
    elif state.has_archive_location == "no":
        query = query.filter(or_(ProjectModel.file_server_location.is_(None), trimmed_archive_root == ""))

    return query.order_by(ranking_score.asc(), normalized_number.asc(), ProjectModel.id.asc())


def _result_from_db_row(row) -> ProjectSearchResult:
    return ProjectSearchResult(
        project=row[0],
        contract_count=int(row[1]),
        recorded_cost_count=int(row[2]),
        original_cost_total=row[3],
        ranking_score=int(row[4]),
        matched_project=bool(row[5]),
        matched_contract=bool(row[6]),
        matched_caan=bool(row[7]),
    )


def search_html(state: ProjectSearchState) -> tuple[list[ProjectSearchResult], bool]:
    """Fetch the bounded HTML result set and whether more matching rows exist."""
    if not state.is_active:
        return [], False
    rows = _ranked_query(state).limit(HTML_RESULT_LIMIT + 1).all()
    has_more = len(rows) > HTML_RESULT_LIMIT
    return [_result_from_db_row(row) for row in rows[:HTML_RESULT_LIMIT]], has_more


def _natural_sort_key(value: str | None):
    return tuple(
        (0, int(chunk)) if chunk.isdigit() else (1, chunk.lower())
        for chunk in re.split(r"(\d+)", value or "")
        if chunk
    )


def _archive_location(project: ProjectModel, user_archives_location: str | None) -> tuple[str, str]:
    """Return public archive path/state without ever falling back to the stored path."""
    root = project.file_server_location.strip() if isinstance(project.file_server_location, str) else None
    if not root:
        return "", "Not recorded"
    if not user_archives_location:
        return "Unavailable", "Recorded"
    try:
        return (
            utils.FileServerUtils.user_path_from_db_data(
                file_server_directories=root,
                user_archives_location=user_archives_location,
            ),
            "Recorded",
        )
    except Exception:
        return "Unavailable", "Recorded"


def _status(value: bool | None, true_label: str, false_label: str) -> str:
    if value is None:
        return "Unknown"
    return true_label if value else false_label


def _safe_cell(value):
    """Neutralize formula text and invalid XML controls in database-backed cells."""
    if not isinstance(value, str):
        return value
    value = ILLEGAL_CHARACTERS_RE.sub("", value)
    if value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


PROJECT_EXPORT_HEADERS = (
    "Result rank", "Project number", "Project name", "Project Information URL",
    "Status", "Drawings", "Campus client",
    "Project manager", "Inspector", "Archive location", "Archive root status",
    "Initial contract value", "Contracts with recorded initial cost", "Linked contracts",
    "Matched in",
)
TRAILING_PROJECT_EXPORT_HEADERS = ("Ranking band", "database index")
CONTRACT_EXPORT_HEADERS = (
    "Contract number", "Contractor", "Executive design organization", "Scope description",
    "Cost estimate", "Original contract cost", "Change-order total",
    "Revised total including change orders", "Funding number", "Bid date", "Contract date",
    "Notice-to-proceed date", "Beneficial occupancy date", "Substantial completion date",
    "Certificate of occupancy date", "Notice of completion date",
    "Notice of completion recorded date", "Termination date", "Current expected end date",
    "Original duration (days)", "Change-order time (days)", "Current duration (days)",
)
CONTRACT_EXPORT_FIELDS = (
    "contract_number", "contractor_org_name", "executive_design_org_name", "scope_description",
    "cost_estimate", "original_contract_cost", "change_order_total", "change_order_revised_cost",
    "funding_number", "bid_date", "contract_date", "ntp_start_date", "beneficial_occupancy_date",
    "substantial_completion_date", "certificate_of_occupancy_date", "noc_completion_date",
    "noc_recorded_date", "termination_date", "change_order_revised_expected_end",
    "original_project_duration", "change_order_time_total", "change_order_revised_duration",
)

PROJECT_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
CONTRACT_HEADER_FILL = PatternFill("solid", fgColor="476A8A")
TRAILING_HEADER_FILL = PatternFill("solid", fgColor="5B4B73")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
DATA_ALIGNMENT = Alignment(vertical="top")
WRAPPED_DATA_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
INFO_LABEL_FILL = PatternFill("solid", fgColor="D9EAF7")
BANDED_ROW_FILL = PatternFill("solid", fgColor="F4F8FC")
THIN_BOTTOM_BORDER = Border(bottom=Side(style="thin", color="D9E2F3"))
URL_FONT = Font(color="0563C1", underline="single")
CURRENCY_HEADERS = frozenset({
    "Cost estimate",
    "Original contract cost",
    "Change-order total",
    "Revised total including change orders",
})
DATE_HEADERS = frozenset({
    "Bid date",
    "Contract date",
    "Notice-to-proceed date",
    "Beneficial occupancy date",
    "Substantial completion date",
    "Certificate of occupancy date",
    "Notice of completion date",
    "Notice of completion recorded date",
    "Termination date",
    "Current expected end date",
})
WRAPPED_HEADERS = frozenset({
    "Project name",
    "Project Information URL",
    "Archive location",
    "Scope description",
    "Funding number",
})
PROJECT_EXPORT_COLUMN_WIDTHS = (
    12, 16, 46, 58, 12, 12, 28, 24, 24, 52, 18, 22, 15, 15, 16,
    18, 30, 34, 60, 17, 18, 18, 22, 22, 16, 16, 18, 21, 23, 26,
    26, 26, 26, 26, 26, 26, 18, 20, 16,
)


def _header_cell(worksheet, value: str, column_index: int):
    cell = WriteOnlyCell(worksheet, value=value)
    if column_index < len(PROJECT_EXPORT_HEADERS):
        cell.fill = PROJECT_HEADER_FILL
    elif column_index < len(PROJECT_EXPORT_HEADERS) + len(CONTRACT_EXPORT_HEADERS):
        cell.fill = CONTRACT_HEADER_FILL
    else:
        cell.fill = TRAILING_HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = HEADER_ALIGNMENT
    cell.border = THIN_BOTTOM_BORDER
    return cell


def _data_cell(worksheet, value, header: str, is_banded_row: bool):
    """Prepare one bounded-memory, presentation-formatted workbook cell."""
    cell = WriteOnlyCell(worksheet, value=_safe_cell(value))
    cell.alignment = WRAPPED_DATA_ALIGNMENT if header in WRAPPED_HEADERS else DATA_ALIGNMENT
    cell.border = THIN_BOTTOM_BORDER
    if is_banded_row:
        cell.fill = BANDED_ROW_FILL
    if header in CURRENCY_HEADERS and isinstance(value, Decimal):
        cell.number_format = '$#,##0.00'
    elif header in DATE_HEADERS and isinstance(value, (date, datetime)):
        cell.number_format = "mmm d, yyyy"
    if header == "Project Information URL" and isinstance(value, str):
        cell.hyperlink = value
        cell.font = URL_FONT
    return cell


def _configure_projects_sheet(worksheet, headers: tuple[str, ...]):
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    worksheet.sheet_view.showGridLines = False
    for column_index, width in enumerate(PROJECT_EXPORT_COLUMN_WIDTHS, start=1):
        worksheet.column_dimensions[get_column_letter(column_index)].width = width
    worksheet.row_dimensions[1].height = 36


def _information_row(worksheet, label: str, value):
    label_cell = WriteOnlyCell(worksheet, value=label)
    label_cell.fill = INFO_LABEL_FILL
    label_cell.font = Font(bold=True, color="1F1F1F")
    label_cell.alignment = DATA_ALIGNMENT
    label_cell.border = THIN_BOTTOM_BORDER

    value_cell = WriteOnlyCell(worksheet, value=_safe_cell(value))
    value_cell.alignment = WRAPPED_DATA_ALIGNMENT
    value_cell.border = THIN_BOTTOM_BORDER
    if label == "Search results URL" and isinstance(value, str):
        value_cell.hyperlink = value
        value_cell.font = URL_FONT
    return label_cell, value_cell


def _configure_information_sheet(worksheet):
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.column_dimensions["A"].width = 30
    worksheet.column_dimensions["B"].width = 100
    worksheet.row_dimensions[1].height = 24


def _project_export_values(result: ProjectSearchResult, rank: int, user_archives_location: str | None):
    project = result.project
    archive_location, archive_status = _archive_location(project, user_archives_location)
    return (
        rank,
        project.number,
        project.name,
        flask.url_for(
            "project_tools.project_info", project_id=project.id, _external=True
        ),
        _status(project.closed, "Closed", "Open"),
        _status(project.drawings, "Yes", "No"),
        project.campus_client or "",
        project.project_manager_name or "",
        project.inspector_name or "",
        archive_location,
        archive_status,
        result.initial_contract_value,
        result.recorded_cost_count,
        result.contract_count,
        result.matched_in,
    )


def _trailing_project_export_values(result: ProjectSearchResult):
    return result.ranking_band, result.project.id


def _contract_export_values(contract: ContractModel | None):
    if contract is None:
        return ("",) * len(CONTRACT_EXPORT_FIELDS)
    return tuple(getattr(contract, field) for field in CONTRACT_EXPORT_FIELDS)


def _chunks(iterator, chunk_size: int):
    while chunk := list(islice(iterator, chunk_size)):
        yield chunk


def build_export_workbook(state: ProjectSearchState, user_archives_location: str | None) -> BytesIO:
    """Build a write-only XLSX workbook for every matching project and contract."""
    if not state.is_active:
        raise ProjectSearchValidationError("Supply a search query or at least one active filter.")

    workbook = Workbook(write_only=True)
    projects_sheet = workbook.create_sheet("Projects and contracts")
    all_export_headers = (
        PROJECT_EXPORT_HEADERS + CONTRACT_EXPORT_HEADERS + TRAILING_PROJECT_EXPORT_HEADERS
    )
    _configure_projects_sheet(projects_sheet, all_export_headers)
    projects_sheet.append(tuple(
        _header_cell(projects_sheet, header, column_index)
        for column_index, header in enumerate(all_export_headers)
    ))
    export_row_count = 0
    project_count = 0

    db_rows = (_result_from_db_row(row) for row in _ranked_query(state).yield_per(100))
    for result_chunk in _chunks(db_rows, 100):
        project_ids = [result.project.id for result in result_chunk]
        contracts_by_project = {project_id: [] for project_id in project_ids}
        contracts = ContractModel.query.filter(ContractModel.project_id.in_(project_ids)).all()
        for contract in contracts:
            contracts_by_project[contract.project_id].append(contract)
        for contract_list in contracts_by_project.values():
            contract_list.sort(key=lambda contract: (_natural_sort_key(contract.contract_number), contract.id))

        for result in result_chunk:
            project_count += 1
            project_values = _project_export_values(
                result, project_count, user_archives_location
            )
            contracts_for_project = contracts_by_project[result.project.id] or [None]
            for contract in contracts_for_project:
                row_values = (
                    project_values
                    + _contract_export_values(contract)
                    + _trailing_project_export_values(result)
                )
                projects_sheet.append(tuple(
                    _data_cell(
                        projects_sheet,
                        value,
                        header,
                        is_banded_row=export_row_count % 2 == 1,
                    )
                    for value, header in zip(row_values, all_export_headers)
                ))
                export_row_count += 1

    information_sheet = workbook.create_sheet("Search information")
    _configure_information_sheet(information_sheet)
    information_sheet.append((
        _header_cell(information_sheet, "Search information", 0),
        _header_cell(information_sheet, "Value", 0),
    ))
    information_sheet.append(_information_row(information_sheet, "Query", state.query))
    search_results_url = flask.url_for(
        "project_tools.project_search",
        _external=True,
        **state.export_parameters(),
    )
    information_sheet.append(_information_row(
        information_sheet, "Search results URL", search_results_url
    ))
    information_sheet.append(_information_row(information_sheet, "Status", state.status))
    information_sheet.append(_information_row(information_sheet, "Drawings", state.drawings))
    information_sheet.append(_information_row(
        information_sheet,
        "File server location",
        {"any": "Any", "yes": "Known", "no": "Unknown"}[state.has_archive_location],
    ))
    information_sheet.append(_information_row(
        information_sheet, "Export timestamp (UTC)", datetime.now(timezone.utc).isoformat()
    ))
    information_sheet.append(_information_row(
        information_sheet, "Total matched projects", project_count
    ))
    information_sheet.append(_information_row(
        information_sheet, "Total exported rows", export_row_count
    ))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
