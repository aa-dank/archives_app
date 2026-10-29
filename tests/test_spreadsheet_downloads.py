from datetime import datetime, timezone
from io import BytesIO
import re

from openpyxl import load_workbook
import pandas as pd
import pytest

from archives_application import create_app
from archives_application.archiver import routes as archive_routes


class TestConfig:
    SECRET_KEY = "spreadsheet-download-test"
    SQLALCHEMY_DATABASE_URI = "sqlite://"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    GOOGLE_CLIENT_ID = "spreadsheet-download-test-client"
    KEY_VALUE_STORE_URL = "redis://localhost:6379/0"


@pytest.mark.parametrize("url,stem,sheets", [
    ("/endpoints_index?spreadsheet=True", "endpoints", ["Endpoints"]),
    (
        "/dir_contents_summary/download?path=Z:/Archives",
        "dir_contents_summary", ["Child Directories", "Files in Current Dir", "All Files Recursive"],
    ),
])
def test_spreadsheet_downloads_include_utc_timestamp(monkeypatch, url, stem, sheets):
    monkeypatch.setattr(
        archive_routes, "_build_dir_contents_summary_df",
        lambda user_path: (
            pd.DataFrame({"Name": ["Child"]}),
            pd.DataFrame({"Name": ["Current file"]}),
            pd.DataFrame({"Name": ["Recursive file"]}), {},
        ),
    )
    app = create_app(TestConfig)
    before = datetime.now(timezone.utc).replace(microsecond=0)
    response = app.test_client().get(url)
    after = datetime.now(timezone.utc)
    assert response.status_code == 200
    match = re.search(
        rf"filename={stem}_(\d{{8}}T\d{{6}}Z)\.xlsx",
        response.headers["Content-Disposition"],
    )
    assert match
    timestamp = datetime.strptime(match[1], "%Y%m%dT%H%M%SZ").replace(tzinfo=timezone.utc)
    assert before <= timestamp <= after
    assert load_workbook(BytesIO(response.data)).sheetnames == sheets
