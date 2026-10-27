import copy
import json
import os
import time
from collections import Counter
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta, timezone
from io import BytesIO
from types import SimpleNamespace
from urllib.parse import parse_qs, urlsplit

import pytest
from flask import Flask
from openpyxl import load_workbook
from sqlalchemy import text

from archives_application import create_app, db
from archives_application import app_config
from archives_application.archiver import archive_search as search
from archives_application.archiver import archive_search_cache as caching
from archives_application.archiver import routes
from archives_application.cache import CacheUnavailable, KeyValueCache, positive_integer
from archives_application.models import ArchiveSearchRunModel, UserModel


class Clock:
    """Control source-generation and backend-expiration instants independently of wall time."""

    def __init__(self):
        """Start every test at a fixed aware UTC instant."""
        self.now = datetime(2026, 10, 13, 12, tzinfo=timezone.utc)

    def advance(self, seconds):
        """Advance simulated time without sleeping or renewing stored entries."""
        self.now += timedelta(seconds=seconds)


class MemoryCache:
    """Deterministic test backend; never used by the application."""

    def __init__(self, clock):
        """Keep disposable entries, operation histories, and outage state for one test."""
        self.clock = clock
        self.entries = {}
        self.writes = []
        self.reads = []
        self.unavailable = False

    def get(self, key):
        """Record lookup and return only entries still valid on the simulated clock."""
        self.reads.append(key)
        if self.unavailable:
            raise CacheUnavailable()
        entry = self.entries.get(key)
        if entry is None or self.clock.now >= entry[1]:
            return None
        return entry[0]

    def set(self, key, value, ttl_seconds):
        """Publish one complete entry with its supplied TTL and record the write."""
        positive_integer(ttl_seconds, "ttl_seconds")
        if self.unavailable:
            raise CacheUnavailable()
        self.entries[key] = (value, self.clock.now + timedelta(seconds=ttl_seconds))
        self.writes.append((key, value, ttl_seconds))

    def delete(self, key):
        """Remove a disposable key without affecting other entries."""
        self.entries.pop(key, None)


class TestConfig:
    """Enable cache policy with SQLite telemetry and a lazy, unused key-value connection."""
    SECRET_KEY = "archive-search-cache-test"
    SQLALCHEMY_DATABASE_URI = "sqlite://"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    GOOGLE_CLIENT_ID = "cache-test-client"
    KEY_VALUE_STORE_URL = "redis://localhost:6379/0"
    WTF_CSRF_ENABLED = False
    ARCHIVE_SEARCH_CACHE_ENABLED = True
    CACHE_NAMESPACE = "test:archives"
    ARCHIVES_LOCATION = "/mnt/archives/Records"
    USER_ARCHIVES_LOCATION = r"N:\PPDO\Records"


@pytest.fixture()
def environment(monkeypatch, tmp_path):
    """Provide real SQLite telemetry, deterministic cache/time, and counted source-query boundaries."""
    clock = Clock()
    backend = MemoryCache(clock)
    monkeypatch.setattr(caching, "utc_now", lambda: clock.now)
    monkeypatch.setattr(search, "utc_now", lambda: clock.now)
    app = create_app(TestConfig)
    assert isinstance(app.extensions["cache"], KeyValueCache)
    app.extensions["cache"] = backend
    monkeypatch.setattr(search.utils.FlaskAppUtils, "create_temp_filepath",
                        lambda filename, **kwargs: str(tmp_path / filename))
    calls = Counter()

    def resolve(request, app):
        """Retain complete scope facts and warnings while recording source resolution work."""
        calls["scope"] += 1
        return search.ScopeResolution(
            request.requested_scope_type, request.requested_scope_value or "",
            prefixes=["13xx   Ranch"] if request.requested_scope_type != "all" else [],
            messages=["Scope note"], warnings=["Coverage warning"],
            project_count=2, usable_project_count=1, missing_project_location_count=1,
            roots_with_no_indexed_files=["Empty root"], caan_found=True, linked_project_count=2,
        )

    def counted(name, value):
        """Build a query stub whose independent results reveal whether warm hits repeat work."""
        def call(*args, **kwargs):
            """Count one source-query stage and return a fresh copy of its fixture data."""
            calls[name] += 1
            return copy.deepcopy(value)
        return call

    monkeypatch.setattr(search, "resolve_scope", resolve)
    monkeypatch.setattr(search, "_execute_content_search", counted("content", [
        {"file_hash": "hash-1", "content_rank": 0.7, "matching_chunks": 2, "best_chunk_id": 1},
    ]))
    monkeypatch.setattr(search, "_execute_filename_search", counted("filename", [
        {"file_hash": "hash-1", "filepath_rank": 0.3, "best_location_id": 1, "matching_location_ids": [1]},
        {"file_hash": "hash-2", "filepath_rank": 0.2, "best_location_id": 3, "matching_location_ids": [3]},
    ]))
    monkeypatch.setattr(search, "_fetch_file_metadata", counted("metadata", {
        key: {"extension": "pdf", "size_bytes": 2048, "has_chunks": True, "text_length": 500}
        for key in ("hash-1", "hash-2")
    }))
    locations = {}
    for key, ids in (("hash-1", (1, 2)), ("hash-2", (3,))):
        locations[key] = [{
            "file_hash": key, "location_id": number, "filename": "Plans   café <x>.pdf",
            "file_server_directories": "13xx   Ranch", "in_scope": True,
            "user_path": r"N:\PPDO\Records\13xx   Ranch\Plans   café <x>.pdf",
            "existence_confirmed": datetime(2026, 10, 1, 10, 30), "hash_confirmed": None,
        } for number in ids]
    monkeypatch.setattr(search, "_fetch_locations", counted("locations", locations))
    monkeypatch.setattr(search, "_fetch_snippets", counted("snippets", {1: "&lt;x&gt; <mark>plans</mark> café"}))
    coverage = {name: 0 for name in caching.COVERAGE_FIELDS}
    coverage.update(files_in_scope=2, content_searchable_files=2, filename_path_searchable_files=2)
    monkeypatch.setattr(search, "_coverage_summary", counted("coverage", coverage))
    with app.app_context():
        UserModel.__table__.create(db.engine)
        db.session.add_all([UserModel(id=1, email="one@example.org"), UserModel(id=2, email="two@example.org")])
        db.session.execute(text("""
            CREATE TABLE archive_search_runs (
                id INTEGER PRIMARY KEY, search_timestamp TIMESTAMP DEFAULT CURRENT_TIMESTAMP NOT NULL,
                duration_ms BIGINT, user_id INTEGER, query_text TEXT NOT NULL, search_mode VARCHAR(32) NOT NULL,
                requested_scope_type VARCHAR(20) NOT NULL, requested_scope_value TEXT,
                extension_filters JSON NOT NULL, status VARCHAR(20) NOT NULL,
                request_source VARCHAR(20) NOT NULL, returned_result_count BIGINT,
                coverage_summary JSON, application_version TEXT
            )
        """))
        db.session.commit()
    env = SimpleNamespace(app=app, client=app.test_client(), clock=clock, backend=backend, calls=calls, path=tmp_path)
    yield env
    with app.app_context():
        db.session.remove()
        db.session.execute(text("DROP TABLE archive_search_runs"))
        UserModel.__table__.drop(db.engine)


def page(env, **options):
    """Prepare a browser page with standard limits and optional server-side cache controls."""
    return search.prepare_archive_search_page(
        search.ArchiveSearchRequest.from_values("plans"), env.app, 300, 3000, None, **options,
    )


def result_entry(env):
    """Find the sole data generation without confusing its separate workbook reference."""
    return next((key, value) for key, value in env.backend.entries.items() if ":result:" in key)


def test_warm_page_skips_all_search_stages_and_excel(environment, monkeypatch):
    """Verify full source reuse and export reuse while every request still persists new telemetry."""
    env = environment
    with env.app.app_context():
        first = page(env)
        monkeypatch.setattr(search, "write_archive_search_workbook", lambda *args: pytest.fail("warm export"))
        second = page(env)
        records = db.session.query(ArchiveSearchRunModel).order_by(ArchiveSearchRunModel.id).all()
        assert len(records) == 2 and records[0].id != records[1].id
        assert all(record.status == "successful" and record.returned_result_count == 2 for record in records)
        assert records[0].coverage_summary == records[1].coverage_summary
    assert first.search == second.search
    assert first.timestamp == second.timestamp
    assert not first.cache_hit and second.cache_hit
    assert first.generated_at == second.generated_at == env.clock.now
    assert first.expires_at == second.expires_at == env.clock.now + timedelta(days=5)
    assert all(count == 1 for count in env.calls.values())
    assert len(env.backend.writes) == 2


def test_alias_refresh_and_result_headers(environment):
    """Verify shared alias identity, preserved active refresh options, escaping, and HTTP freshness policy."""
    env = environment
    params = {"search_term": ' "plans"   café! ', "search_mode": "filepath", "scope_type": "location",
              "location_scope": r"N:\PPDO\Records\13xx   Ranch", "file_extension": ".PDF, DOCX",
              "project_number": "", "caan": ""}
    first = env.client.get("/archives_search", query_string=params)
    cached = env.client.get("/file_search", query_string=params)
    assert first.status_code == cached.status_code == 200
    assert first.headers["Cache-Control"] == cached.headers["Cache-Control"] == "private, no-store"
    assert b"archive-cache-state-fresh\">Fresh results</span>" in first.data
    assert b"archive-cache-state-cached\">Cached results</span>" in cached.data
    assert b'Coverage &amp; freshness' in cached.data
    assert b'<time datetime="2026-10-13T12:00:00+00:00">2026-10-13 12:00:00 UTC</time>' in cached.data
    assert b'<time datetime="2026-10-18T12:00:00+00:00">2026-10-18 12:00:00 UTC</time>' in cached.data
    assert "13xx   Ranch" in cached.text and "Plans   café &lt;x&gt;.pdf" in cached.text
    import html
    import re
    link = html.unescape(re.search(r'href="([^"]+)">Run fresh search', cached.text)[1])
    active = parse_qs(urlsplit(link).query)
    assert active["refresh"] == ["1"]
    assert active["search_term"] == ['"plans"   café!']
    for key in ("search_mode", "scope_type", "location_scope", "file_extension"):
        assert active[key] == [params[key]]
    assert "project_number" not in active and "caan" not in active
    env.clock.advance(10)
    assert env.client.get(link).status_code == 200
    assert env.calls["scope"] == 2
    assert env.client.get("/archives_search", query_string=params).status_code == 200
    assert env.calls["scope"] == 2


@pytest.mark.parametrize("query", [
    "refresh=1", "search_term=&refresh=1", "search_term=x&refresh=", "search_term=x&refresh=0",
    "search_term=x&refresh=true", "search_term=x&refresh=1&refresh=1", "search_term=x&unknown=y",
    "search_term=x&search_term=y", "timestamp=20261013120000&search_term=x",
    "timestamp=20261013120000&refresh=1", "timestamp=20269999120000", "timestamp=../../etc/passwd",
    "timestamp=20261013120000&timestamp=20261013120000",
])
def test_invalid_requests_never_reach_cache(environment, query):
    """Reject ambiguous browser/download controls before source execution or cache access."""
    env = environment
    assert env.client.get("/archives_search?" + query).status_code == 400
    assert not env.calls and not env.backend.reads and not env.backend.writes


def test_api_is_fresh_and_metadata_stays_out_of_response(environment, monkeypatch):
    """Keep authenticated API requests fresh and preserve their transport contract despite a warm browser cache."""
    env = environment
    monkeypatch.setattr(routes, "_archive_search_api_user", lambda payload: SimpleNamespace(id=1))
    assert env.client.get("/archives_search?search_term=plans").status_code == 200
    for _ in range(2):
        response = env.client.post("/api/archives_search", json={"query_text": "plans"})
        assert response.status_code == 200
        assert not {"cache_hit", "generated_at", "expires_at", "generation_id"} & response.json.keys()
    assert env.calls["scope"] == 3
    with env.app.app_context():
        assert [row.request_source for row in db.session.query(ArchiveSearchRunModel).order_by(ArchiveSearchRunModel.id)] == ["web", "api", "api"]
    monkeypatch.setattr(routes, "_archive_search_api_user", lambda payload: None)
    assert env.client.post("/api/archives_search", json={"query_text": "plans"}).status_code == 401


def test_current_viewer_telemetry_and_telemetry_outage_on_hit(environment, monkeypatch):
    """Attribute each hit to its current viewer and keep telemetry outages nonfatal."""
    env = environment
    with env.app.app_context():
        request = search.ArchiveSearchRequest.from_values("plans")
        policy = caching.browser_cache_policy(env.app)
        first = search.ArchiveSearchRun(request, env.app, 3000, user_id=1, cache_policy=policy)
        first.execute()
        second = search.ArchiveSearchRun(request, env.app, 3000, user_id=2, cache_policy=policy)
        second.execute()
        assert second.generation.cache_hit
        rows = db.session.query(ArchiveSearchRunModel).order_by(ArchiveSearchRunModel.id).all()
        assert [row.user_id for row in rows] == [1, 2]
        assert first.record_id != second.record_id
        def fail_telemetry_flush():
            """Simulate persistence failure without changing the cached search source."""
            raise RuntimeError("telemetry outage")

        monkeypatch.setattr(db.session, "flush", fail_telemetry_flush)
        third = search.ArchiveSearchRun(request, env.app, 3000, cache_policy=policy)
        assert third.execute() == second.search_data
        assert third.record_id is None and third.generation.cache_hit


def test_expiration_and_missing_workbook_do_not_renew_results(environment):
    """Regenerate only a missing workbook, preserve source expiry, and retain downloads after result expiration."""
    env = environment
    with env.app.app_context():
        first = page(env)
        key, (payload, backend_expiry) = result_entry(env)
        env.clock.advance(4 * 86400)
        os.unlink(env.path / f"{search.WORKBOOK_FILENAME_PREFIX}{first.timestamp}.xlsx")
        regenerated = page(env)
        assert regenerated.cache_hit and regenerated.timestamp != first.timestamp
        assert regenerated.generated_at == first.generated_at and regenerated.expires_at == first.expires_at
        assert env.backend.entries[key] == (payload, backend_expiry)
        assert env.backend.writes[-1][2] == 86400
        workbook = load_workbook(env.path / f"{search.WORKBOOK_FILENAME_PREFIX}{regenerated.timestamp}.xlsx")
        coverage = dict(workbook["Coverage"].values)
        assert coverage["generated_at (UTC)"] == "2026-10-13 12:00:00 UTC"
        env.clock.advance(86400)
        fresh = page(env)
        assert not fresh.cache_hit and env.calls["scope"] == 2
    retained = env.client.get("/archives_search", query_string={"timestamp": regenerated.timestamp})
    assert retained.status_code == 200 and retained.data.startswith(b"PK")
    assert retained.headers.get("Cache-Control") != "private, no-store"


def test_failed_refresh_preserves_old_entry(environment, monkeypatch):
    """Propagate a refresh failure while leaving the prior generation usable for normal requests."""
    env = environment
    with env.app.app_context():
        first = page(env)
        previous = copy.deepcopy(env.backend.entries)
        original = search.ArchiveSearchRun._execute_search
        def fail_search(self):
            """Fail explicit refresh before it can publish a replacement generation."""
            raise RuntimeError("search failed")

        monkeypatch.setattr(search.ArchiveSearchRun, "_execute_search", fail_search)
        with pytest.raises(RuntimeError, match="search failed"):
            page(env, refresh=True)
        assert env.backend.entries == previous
        monkeypatch.setattr(search.ArchiveSearchRun, "_execute_search", original)
        assert page(env).timestamp == first.timestamp
        statuses = [row.status for row in db.session.query(ArchiveSearchRunModel).order_by(ArchiveSearchRunModel.id)]
        assert statuses == ["successful", "failed", "successful"]


def test_empty_and_warning_results_cache(environment, monkeypatch):
    """Reuse successful empty/warning results with the same fixed lifetime as populated results."""
    env = environment
    monkeypatch.setattr(search, "_execute_content_search", lambda *args: [])
    monkeypatch.setattr(search, "_execute_filename_search", lambda *args: [])
    with env.app.app_context():
        first = page(env)
        second = page(env)
        assert first.search["results"] == [] and second.cache_hit
        assert second.search["warnings"] == ["Coverage warning"]
        assert first.expires_at - first.generated_at == timedelta(days=5)


def test_location_root_resolving_to_all_archives_caches(environment, monkeypatch):
    """Accept the existing Records-root scope conversion without losing cached coverage warnings."""
    env = environment
    monkeypatch.setattr(search, "resolve_scope", lambda request, app: search.ScopeResolution(
        "all", request.requested_scope_value, warnings=["Location resolves to Records root"],
    ))
    request = search.ArchiveSearchRequest.from_values("plans", requested_scope_type="location",
                                                      requested_scope_value=env.app.config["USER_ARCHIVES_LOCATION"])
    with env.app.app_context():
        first = search.prepare_archive_search_page(request, env.app, 300, 3000, None)
        second = search.prepare_archive_search_page(request, env.app, 300, 3000, None)
    assert second.cache_hit and second.search == first.search
    assert second.search["scope"].scope_type == "all"


@pytest.mark.parametrize("damage", ["missing", "json", "kind", "timestamp", "extra"])
def test_invalid_workbook_references_regenerate_only_artifact(environment, damage):
    """Treat missing or malformed artifact references as export misses without repeating source queries."""
    env = environment
    with env.app.app_context():
        first = page(env)
        original_result = copy.deepcopy(result_entry(env))
        key = next(key for key in env.backend.entries if ":workbook:" in key)
        payload, expires = env.backend.entries[key]
        reference = json.loads(payload)
        if damage == "missing": env.backend.delete(key)
        else:
            if damage == "kind": reference["storage_kind"] = "other"
            if damage == "timestamp": reference["timestamp"] = "../../secret"
            if damage == "extra": reference["path"] = "/secret"
            payload = b"invalid" if damage == "json" else json.dumps(reference).encode()
            env.backend.entries[key] = (payload, expires)
        env.clock.advance(5)
        second = page(env)
        assert second.cache_hit and second.export_available and second.timestamp != first.timestamp
        assert second.generated_at == first.generated_at and second.expires_at == first.expires_at
        assert env.calls["scope"] == 1 and result_entry(env) == original_result


@pytest.mark.parametrize("operation", ["result", "reference"])
def test_failed_writes_preserve_current_fresh_workbook(environment, monkeypatch, operation):
    """Keep current exports usable when result or artifact-reference storage fails."""
    env = environment
    original_set = env.backend.set
    def fail_selected(key, *args):
        """Fail only the selected storage layer to isolate data-publication and artifact-publication fallback."""
        if (":result:" if operation == "result" else ":workbook:") in key:
            raise CacheUnavailable()
        original_set(key, *args)
    monkeypatch.setattr(env.backend, "set", fail_selected)
    with env.app.app_context():
        first = page(env)
        assert first.export_available and not first.cache_hit
        if operation == "result":
            assert first.expires_at is None and not env.backend.entries
        else:
            assert page(env).cache_hit
            assert len([key for key in env.backend.entries if ":workbook:" in key]) == 0


def test_search_duration_excludes_workbook_and_render_time(environment, monkeypatch):
    """Ensure telemetry measures source/cache work while expensive export work stays outside the timer."""
    env = environment
    counter = [0.0]
    monkeypatch.setattr(search.time, "perf_counter", lambda: counter[0])
    original_search = search.ArchiveSearchRun._execute_search
    original_writer = search.write_archive_search_workbook
    def timed_search(self):
        """Charge simulated monotonic time to the source phase measured by telemetry."""
        counter[0] += 2
        return original_search(self)
    def timed_workbook(*args):
        """Charge a larger export cost that must not appear in search duration."""
        counter[0] += 10
        return original_writer(*args)
    monkeypatch.setattr(search.ArchiveSearchRun, "_execute_search", timed_search)
    monkeypatch.setattr(search, "write_archive_search_workbook", timed_workbook)
    with env.app.app_context():
        page(env)
        page(env)
        durations = [row.duration_ms for row in db.session.query(ArchiveSearchRunModel).order_by(ArchiveSearchRunModel.id)]
    assert durations == [2000, 0]


def test_cache_logs_omit_query_paths_and_credentials(environment, caplog):
    """Expose cache outcomes and resource counters without recording sensitive source inputs."""
    env = environment
    caplog.set_level("INFO")
    with env.app.app_context():
        request = search.ArchiveSearchRequest.from_values("private query", requested_scope_type="location",
                                                          requested_scope_value="private path")
        search.prepare_archive_search_page(request, env.app, 300, 3000, None)
        search.prepare_archive_search_page(request, env.app, 300, 3000, None)
    logs = "\n".join(record.message for record in caplog.records if "archive_search_cache" in record.message)
    assert "private query" not in logs and "private path" not in logs
    assert "lookup_hit" in logs and "workbook_reuse" in logs and "write_success" in logs
    assert "identity_digest" in logs and "payload_bytes" in logs and "source_age_seconds" in logs


def test_html_limit_slices_independently(environment):
    """Apply per-response display limits after cache lookup without truncating cached export data."""
    env = environment
    with env.app.app_context():
        first = page(env)
        second = search.prepare_archive_search_page(search.ArchiveSearchRequest.from_values("plans"), env.app, 1, 3000, None)
        assert second.cache_hit and len(second.search["results"]) == 1
        assert second.search["html_limit_hit"]
        assert len(first.search["results"]) == len(page(env).search["results"]) == 2


@pytest.mark.parametrize("setting,value", [
    ("ARCHIVE_SEARCH_CACHE_ENABLED", False), ("ARCHIVE_SEARCH_CACHE_ENABLED", "False"),
    ("ARCHIVE_SEARCH_CACHE_TTL_SECONDS", 0), ("ARCHIVE_SEARCH_CACHE_TTL_SECONDS", True),
    ("ARCHIVE_SEARCH_WORKBOOK_RETENTION_SECONDS", 432000), ("CACHE_MAX_ENTRY_BYTES", 10),
])
def test_disabled_invalid_and_oversized_fall_back(environment, setting, value):
    """Keep complete fresh results available when policy or entry size prevents caching."""
    env = environment
    env.app.config[setting] = value
    with env.app.app_context():
        assert not page(env).cache_hit
        assert not page(env).cache_hit
    assert env.calls["scope"] == 2 and not env.backend.entries


def test_unavailable_cache_is_bounded_and_fresh(environment):
    """Pay one failed backend read per request and preserve fresh source/export behavior."""
    env = environment
    env.backend.unavailable = True
    with env.app.app_context():
        first = page(env)
        second = page(env)
    assert first.export_available and second.export_available
    assert env.calls["scope"] == 2 and len(env.backend.reads) == 2
    assert not env.backend.writes


def test_entry_expiring_during_decode_runs_fresh(environment, monkeypatch):
    """A pre-decode freshness check cannot authorize a hit after decoding crosses expiry."""
    env = environment
    with env.app.app_context():
        first = page(env)
        env.clock.advance(432000 - 1)
        original_decode = caching.decode_search_data
        delayed_once = False

        def decode_across_expiration(*args):
            """Let the first hit's full reconstruction cross its source expiration."""
            nonlocal delayed_once
            result = original_decode(*args)
            if not delayed_once:
                env.clock.advance(2)
                delayed_once = True
            return result

        monkeypatch.setattr(caching, "decode_search_data", decode_across_expiration)
        fresh = page(env)
        assert not fresh.cache_hit
        assert fresh.generated_at > first.expires_at
        assert fresh.export_available and env.calls["scope"] == 2


def test_unrepresentable_expiration_preserves_fresh_search(environment):
    """A positive lifetime beyond datetime range fails only cache publication."""
    env = environment
    oversized_lifetime = 10 ** 30
    env.app.config["ARCHIVE_SEARCH_WORKBOOK_RETENTION_SECONDS"] = oversized_lifetime + 1
    with env.app.app_context():
        fresh = page(env, cache_ttl_seconds=oversized_lifetime)
        assert not fresh.cache_hit and fresh.expires_at is None
        assert fresh.export_available and len(fresh.search["results"]) == 2
        assert not env.backend.entries
        assert db.session.query(ArchiveSearchRunModel).one().status == "successful"


@pytest.mark.parametrize("operation", ["get", "set", "delete"])
def test_backend_connection_option_overflow_is_cache_unavailable(operation):
    """Overflow in lazy socket/configuration handling stays inside the cache failure boundary."""
    def overflow(*args, **kwargs):
        """Simulate an OS timeout conversion outside the representable numeric range."""
        raise OverflowError("socket timeout cannot be represented")

    adapter = KeyValueCache(SimpleNamespace(get=overflow, set=overflow, delete=overflow))
    with pytest.raises(CacheUnavailable):
        if operation == "set":
            adapter.set("key", b"value", 1)
        else:
            getattr(adapter, operation)("key")


@pytest.mark.parametrize("damage", ["json", "version", "identity", "expired", "future", "type", "bounds", "size", "timestamp", "incomplete"])
def test_invalid_payload_is_a_miss(environment, damage):
    """Reject incompatible, stale, malformed, and oversized wire data before using it as source results."""
    env = environment
    with env.app.app_context():
        page(env)
        key, (payload, expiry) = result_entry(env)
        wire = json.loads(payload)
        if damage == "version":
            wire["schema_version"] = 999
        if damage == "identity":
            wire["identity_digest"] = "wrong"
        if damage == "expired":
            wire["generated_at"] = (env.clock.now - timedelta(days=6)).isoformat()
            wire["expires_at"] = (env.clock.now - timedelta(days=1)).isoformat()
        if damage == "future":
            wire["generated_at"] = (env.clock.now + timedelta(days=1)).isoformat()
        if damage == "type":
            wire["search_data"]["coverage"]["files_in_scope"] = True
        if damage == "bounds":
            wire["search_data"]["results"] *= 1501
        if damage == "timestamp":
            wire["generated_at"] = "2026-10-13T12:00:00"
        if damage == "incomplete":
            del wire["search_data"]["scope"]["prefixes"]
        broken = json.dumps(wire).encode()
        if damage == "json":
            broken = b"not json"
        if damage == "size":
            broken = b"x" * (env.app.config["CACHE_MAX_ENTRY_BYTES"] + 1)
        env.backend.entries[key] = (broken, expiry)
        assert not page(env).cache_hit
    assert env.calls["scope"] == 2


def test_json_roundtrip_and_response_independence(environment):
    """Preserve runtime types, Unicode, spaces, snippets, and independent nested response data."""
    env = environment
    with env.app.app_context():
        run = search.ArchiveSearchRun(search.ArchiveSearchRequest.from_values("plans"), env.app, 3000)
        data = run.execute()
        data["results"][0]["locations"][1]["existence_confirmed"] = datetime(2026, 10, 2, tzinfo=timezone(timedelta(hours=-7)))
        wire = caching.encode_search_data(data, 3000)
        decoded = caching.decode_search_data(json.loads(caching._json_bytes(wire)), 3000)
        assert decoded == data
        assert decoded["results"][0]["matching_location_ids"] == {1}
        dates = decoded["results"][0]["locations"]
        assert dates[0]["existence_confirmed"].tzinfo is None
        assert dates[1]["existence_confirmed"].utcoffset() == timedelta(hours=-7)
        assert "   " in dates[0]["user_path"]
        assert "<mark>" in decoded["results"][0]["snippet"]
        decoded["scope"].warnings.append("later")
        assert data["scope"].warnings == ["Coverage warning"]
        first = page(env)
        first.search["warnings"].append("response warning")
        first.search["results"][0]["locations"].clear()
        second = page(env)
        assert second.search["warnings"] == ["Coverage warning"]
        assert len(second.search["results"][0]["locations"]) == 2


def test_export_failure_does_not_poison_cache(environment, monkeypatch):
    """Keep a transient export warning out of the source generation and recover on a later hit."""
    env = environment
    real_writer = search.pd.ExcelWriter
    def fail_excel_writer(*args, **kwargs):
        """Fail before publishing any workbook bytes or reference."""
        raise OSError("export failed")

    monkeypatch.setattr(search.pd, "ExcelWriter", fail_excel_writer)
    with env.app.app_context():
        failed = page(env)
        assert not failed.export_available and failed.timestamp is None
        assert any("Excel export was unavailable" in item for item in failed.search["warnings"])
        assert not list(env.path.iterdir())
        monkeypatch.setattr(search.pd, "ExcelWriter", real_writer)
        recovered = page(env)
        assert recovered.cache_hit and recovered.export_available
        assert recovered.search["warnings"] == ["Coverage warning"]
        assert env.calls["scope"] == 1


def test_per_operation_ttl_and_delayed_publication(environment, monkeypatch):
    """Isolate lifetime policies and cap delayed result/reference writes at the original expiration."""
    env = environment
    with env.app.app_context():
        default = page(env)
        short = page(env, cache_ttl_seconds=60)
        assert short.expires_at - short.generated_at == timedelta(seconds=60)
        assert short.timestamp != default.timestamp
        env.clock.advance(60)
        assert page(env).cache_hit
        assert not page(env, cache_ttl_seconds=60).cache_hit
        policy = caching.browser_cache_policy(env.app, ttl_seconds=100)
        cache = caching.ArchiveSearchCache(policy, search.ArchiveSearchRequest.from_values("plans"), env.app, 3000)
        generation = caching.fresh_generation()
        real_encode = caching.encode_search_data
        def delayed(*args):
            """Simulate serialization latency before publication computes the remaining backend TTL."""
            env.clock.advance(10)
            return real_encode(*args)
        monkeypatch.setattr(caching, "encode_search_data", delayed)
        source = search.ArchiveSearchRun(search.ArchiveSearchRequest.from_values("plans"), env.app, 3000).execute()
        generation = caching.fresh_generation()
        published = cache.publish(source, generation)
        assert env.backend.writes[-1][2] == 90
        env.clock.advance(90)
        count = len(env.backend.writes)
        cache.publish_workbook(published, short.timestamp)
        assert len(env.backend.writes) == count


@pytest.mark.parametrize("value", [False, True, 0, -1, "60", "bad", 0.5, None])
def test_positive_lifetimes_are_strict(value):
    """Reject bools, coercible strings, fractions, and nonpositive values as cache lifetimes."""
    with pytest.raises(ValueError):
        caching.ArchiveSearchCachePolicy(ttl_seconds=value)


def test_long_override_disables_caching_but_search_works(environment):
    """Disable reuse when a server-side override no longer fits workbook retention."""
    env = environment
    with env.app.app_context():
        assert page(env, cache_ttl_seconds=518400).expires_at is None
        assert not env.backend.entries


def test_identity_inputs_and_safe_normalization(environment, monkeypatch):
    """Isolate source-selection policies while excluding credentials, viewer state, and display-only limits."""
    env = environment
    request = search.ArchiveSearchRequest.from_values('  "Plans"   café! ', "FILEPATH", "project", "ab12", ".PDF, DOCX")
    assert request.query_text == '"Plans"   café!' and request.requested_scope_value == "AB12"
    baseline = caching.identity_digest(request, env.app.config, 3000, 432000)
    from dataclasses import replace
    variants = [replace(request, query_text='"plans"   café!'), replace(request, query_text='"Plans" café!'),
                replace(request, search_mode="combined"), replace(request, requested_scope_type="caan"),
                replace(request, requested_scope_value="AB13"), replace(request, extensions=("docx", "pdf"))]
    for changed in variants:
        assert caching.identity_digest(changed, env.app.config, 3000, 432000) != baseline
    for key, value in {
        "ARCHIVE_SEARCH_CHUNK_CANDIDATE_LIMIT": 40000, "ARCHIVE_SEARCH_CHUNK_CANDIDATE_MULTIPLIER": 10,
        "ARCHIVES_LOCATION": "/other", "USER_ARCHIVES_LOCATION": r"Z:\Other",
        "SQLALCHEMY_DATABASE_URI": "sqlite:///other.db",
    }.items():
        config = dict(env.app.config, **{key: value})
        assert caching.identity_digest(request, config, 3000, 432000) != baseline
    assert caching.identity_digest(request, env.app.config, 100, 432000) != baseline
    assert caching.identity_digest(request, env.app.config, 3000, 60) != baseline
    assert caching.identity_digest(request, dict(env.app.config, ARCHIVE_SEARCH_HTML_LIMIT=1), 3000, 432000) == baseline
    config = dict(env.app.config, SQLALCHEMY_DATABASE_URI="postgresql://first:secret@host:5432/archives")
    other_creds = dict(config, SQLALCHEMY_DATABASE_URI="postgresql://second:other@host/archives")
    assert caching.identity_digest(request, config, 3000, 432000) == caching.identity_digest(request, other_creds, 3000, 432000)
    for url in ("postgresql://host2/archives", "postgresql://host:5433/archives", "postgresql://host/other"):
        assert caching.identity_digest(request, dict(config, SQLALCHEMY_DATABASE_URI=url), 3000, 432000) != caching.identity_digest(request, config, 3000, 432000)
    monkeypatch.setattr(caching, "SEARCH_FORMAT_VERSION", 2)
    assert caching.identity_digest(request, env.app.config, 3000, 432000) != baseline


def test_namespaces_isolate_and_old_artifact_cannot_replace_refresh(environment):
    """Prevent older artifact work from replacing newer results and separate dataset namespaces."""
    env = environment
    with env.app.app_context():
        request = search.ArchiveSearchRequest.from_values("plans")
        policy = caching.browser_cache_policy(env.app)
        old = search.ArchiveSearchRun(request, env.app, 3000, cache_policy=policy)
        old.execute()
        new = page(env, refresh=True)
        before = result_entry(env)
        old.cache.publish_workbook(old.generation, "20261013120000000000")
        assert result_entry(env) == before
        assert page(env).generated_at == new.generated_at
        env.app.config["CACHE_NAMESPACE"] = "test:other"
        assert not page(env).cache_hit
        assert len([key for key in env.backend.entries if ":result:" in key]) == 2


def test_workbook_publication_is_atomic_and_collision_safe(environment, monkeypatch):
    """Publish complete concurrent XLSX files with distinct IDs resolvable by the download route."""
    env = environment
    real_replace = search.os.replace
    observations = []
    def observe(source, destination):
        """Verify the destination is unpublished and staging bytes form a complete workbook before rename."""
        assert not os.path.exists(destination)
        assert load_workbook(source).sheetnames == ["Results", "Locations", "Coverage"]
        observations.append(destination)
        real_replace(source, destination)
    monkeypatch.setattr(search.os, "replace", observe)
    with env.app.app_context():
        run = search.ArchiveSearchRun(search.ArchiveSearchRequest.from_values("plans"), env.app, 3000)
        data = run.execute()
        with ThreadPoolExecutor(max_workers=2) as pool:
            ids = list(pool.map(lambda _: search.write_archive_search_workbook(copy.deepcopy(data), env.clock.now, env.app), range(2)))
    assert len(set(ids)) == 2
    assert len(observations) == 2
    assert all((env.path / f"{search.WORKBOOK_FILENAME_PREFIX}{timestamp}.xlsx").is_file() for timestamp in ids)
    assert not list(env.path.glob("*.reserve")) and not list(env.path.glob(".archive-search-*"))


@pytest.mark.parametrize("setting,value", [
    ("ARCHIVE_SEARCH_CACHE_ENABLED", "False"), ("KEY_VALUE_STORE_URL", None), ("KEY_VALUE_STORE_URL", "invalid"),
    ("KEY_VALUE_STORE_URL", "redis://localhost:6379/1"), ("CACHE_NAMESPACE", " "),
    ("KEY_VALUE_STORE_URL", "redis://localhost:6379/0?unsupported=yes"),
    ("KEY_VALUE_STORE_URL", "rediss://localhost:6379/0?ssl_cert_reqs=invalid"),
    ("KEY_VALUE_STORE_URL", "redis://localhost:6379/0?cache_config=1"),
    ("KEY_VALUE_STORE_URL", "redis://localhost:6379/0?cache=1"),
    ("CACHE_TIMEOUT_SECONDS", False), ("CACHE_TIMEOUT_SECONDS", 0), ("CACHE_TIMEOUT_SECONDS", float("inf")),
    ("CACHE_MAX_ENTRY_BYTES", True), ("ARCHIVE_SEARCH_WORKBOOK_RETENTION_SECONDS", 432000),
])
def test_startup_invalid_configuration_disables_without_connecting(setting, value, caplog):
    """Disable malformed cache configuration without attempting a backend connection."""
    config = type("InvalidConfig", (TestConfig,), {setting: value})
    # An invalid shared URL can also invalidate RQ configuration; isolate the
    # cache's best-effort boundary from the producer's independent validation.
    app = Flask(__name__)
    app.config.from_object(config)
    app_config.configure_cache(app)
    assert app.extensions["cache"] is None
    assert "state=disabled" in caplog.text and "invalid_configuration" in caplog.text


def test_backend_timeout_bytes_retry_and_error_contract():
    """Enforce bounded byte operations and preserve the distinction between backend and caller errors."""
    adapter = KeyValueCache.from_url("redis://localhost:6379/0?socket_timeout=99&decode_responses=True", 0.5)
    options = adapter.client.connection_pool.connection_kwargs
    assert options["socket_timeout"] == options["socket_connect_timeout"] == 0.5
    assert options["decode_responses"] is False
    import redis
    attempts = []

    def fail_once():
        """Count attempts so the no-retry contract works across redis-py versions."""
        attempts.append(True)
        raise redis.ConnectionError("unavailable")

    with pytest.raises(redis.ConnectionError):
        options["retry"].call_with_retry(fail_once, lambda error: None)
    assert len(attempts) == 1
    class BrokenClient:
        """Fail backend operations while retaining the adapter's public input checks."""

        def get(self, key):
            """Simulate a connection error carrying sensitive backend details."""
            raise redis.ConnectionError("credentials must not escape")

        def set(self, *args, **kwargs):
            """Simulate a write timeout after caller input validation."""
            raise redis.TimeoutError()

        def delete(self, key):
            """Simulate an unavailable backend during deletion."""
            raise redis.ConnectionError()
    broken = KeyValueCache(BrokenClient())
    with pytest.raises(CacheUnavailable):
        broken.get("key")
    with pytest.raises(CacheUnavailable):
        broken.set("key", b"value", 1)
    with pytest.raises(CacheUnavailable):
        broken.delete("key")
    with pytest.raises(ValueError):
        broken.set("key", b"value", True)


@pytest.mark.parametrize("retention", [518400, 0, "518400", True])
def test_maintenance_separate_retention(environment, monkeypatch, retention):
    """Retain search workbooks independently while ordinary-file cleanup and task status still complete."""
    from archives_application.main import main_tasks
    env = environment
    env.app.config["ARCHIVE_SEARCH_WORKBOOK_RETENTION_SECONDS"] = retention
    monkeypatch.setattr(main_tasks, "app", env.app)
    records = []
    monkeypatch.setattr(main_tasks.utils.RQTaskUtils, "initiate_task_subroutine", lambda **kwargs: None)
    monkeypatch.setattr(main_tasks.utils.RQTaskUtils, "complete_task_subroutine", lambda **kwargs: records.append(kwargs))
    now = time.time()
    for name, age in {
        "archive_search_results_20261010120000.xlsx": 3 * 86400,
        "archive_search_results_20261008120000000000.xlsx": 5 * 86400,
        "archive_search_results_20261007120000000000.xlsx": 6 * 86400 + 1,
        "preview.png": 4 * 86400,
        "archive_search_results_invalid.xlsx": 4 * 86400,
    }.items():
        path = env.path / name
        path.write_bytes(b"data")
        os.utime(path, (now - age, now - age))
    # Linux ctime cannot be set; keep the old policy and simulate its old files.
    original_ctime = main_tasks.os.path.getctime
    monkeypatch.setattr(main_tasks.os.path, "getctime", lambda path: now - 4 * 86400 if str(path).endswith(("preview.png", "invalid.xlsx")) else original_ctime(path))
    custodian = main_tasks.AppCustodian(3, 2, {})
    custodian.temporary_files_location = str(env.path)
    result = custodian._temp_file_clean_up_task("test-task")
    assert not (env.path / "preview.png").exists()
    assert not (env.path / "archive_search_results_invalid.xlsx").exists()
    assert (env.path / "archive_search_results_20261010120000.xlsx").exists()
    assert (env.path / "archive_search_results_20261008120000000000.xlsx").exists()
    assert (env.path / "archive_search_results_20261007120000000000.xlsx").exists() == (retention != 518400)
    assert records[0]["q_id"] == "test-task" and result["files_removed"] == (3 if retention == 518400 else 2)
