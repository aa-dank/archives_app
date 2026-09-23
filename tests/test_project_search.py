from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
import pytest

from archives_application import create_app, db
from archives_application.models import CAANModel, ContractModel, ProjectCaanModel, ProjectModel


class TestConfig:
    SECRET_KEY = "project-search-test"
    SQLALCHEMY_DATABASE_URI = "sqlite://"
    SQLALCHEMY_TRACK_MODIFICATIONS = False
    GOOGLE_CLIENT_ID = "project-search-test-client"
    REDIS_URL = "redis://localhost:6379/0"
    USER_ARCHIVES_LOCATION = r"Z:\Archives"


@pytest.fixture()
def app():
    app = create_app(TestConfig)
    with app.app_context():
        db.engine.raw_connection().create_function(
            "btrim",
            -1,
            lambda value, characters=None: value.strip(characters) if value else value,
            deterministic=True,
        )
        ProjectModel.__table__.create(db.engine)
        CAANModel.__table__.create(db.engine)
        ProjectCaanModel.__table__.create(db.engine)
        ContractModel.__table__.create(db.engine)

        first = ProjectModel(
            id=1, number="1000", name="=Formula project", closed=False, drawings=True,
            file_server_location="1000/Formula", project_manager_name="Alice",
            campus_client="Campus-only search value",
        )
        duplicate = ProjectModel(
            id=2, number="1000", name="Library remodel", closed=True, drawings=False,
            file_server_location="   ",
        )
        distributed = ProjectModel(
            id=3, number="2000", name="Library project", closed=False, drawings=None,
        )
        db.session.add_all((first, duplicate, distributed))
        db.session.add_all((
            CAANModel(id=1, caan="70A"),
            CAANModel(id=2, caan="7110"),
        ))
        db.session.add_all((
            ProjectCaanModel(project_id=1, caan_id=1),
            ProjectCaanModel(project_id=3, caan_id=2),
        ))
        db.session.add_all((
            ContractModel(
                id=1, project_id=2, contract_number="A-10",
                scope_description="Library construction", original_contract_cost=Decimal("10.00"),
            ),
            ContractModel(
                id=2, project_id=3, contract_number="A-10",
                scope_description="Windows Construction", original_contract_cost=Decimal("20.00"),
                funding_number="18319\nWO00423905",
            ),
            ContractModel(
                id=3, project_id=3, contract_number="A-2", original_contract_cost=None),
        ))
        db.session.commit()
        yield app
        db.session.remove()
        ContractModel.__table__.drop(db.engine)
        ProjectCaanModel.__table__.drop(db.engine)
        CAANModel.__table__.drop(db.engine)
        ProjectModel.__table__.drop(db.engine)


@pytest.fixture()
def client(app):
    return app.test_client()


def test_empty_form_and_public_parameter_validation(client):
    assert client.get("/project_search").status_code == 200
    assert client.get("/project_search?caan=7199").status_code == 400
    assert client.get("/project_search?query=x&query=y").status_code == 400
    assert client.get("/project_search/export").status_code == 400


def test_duplicate_number_results_use_canonical_project_ids(client):
    response = client.get("/project_search?query=%20%201000%20")
    assert response.status_code == 200
    assert b"/project_info?project_id=1" in response.data
    assert b"/project_info?project_id=2" in response.data
    assert b"Download all results" in response.data


def test_distributed_project_and_contract_match_ranks_after_one_contract(app):
    with app.app_context():
        from werkzeug.datastructures import MultiDict
        from archives_application.project_tools.project_search import parse_request, search_html

        results, has_more = search_html(parse_request(MultiDict([("query", "library construction")])) )
        assert not has_more
        assert [result.project.id for result in results] == [2, 3]
        assert results[0].ranking_band == "One linked contract"
        assert results[1].ranking_band == "Distributed metadata match"
        assert results[1].matched_in == "Project, Contract"


def test_filter_only_search_and_blank_root_handling(client):
    response = client.get("/project_search?status=closed&has_archive_location=no")
    assert response.status_code == 200
    assert b"Library remodel" in response.data
    assert b"File server location" in response.data
    assert b"Archive root" not in response.data
    assert b"File server location: Unknown" in response.data
    assert b"Matched in</th>" not in response.data


def test_drawings_yes_or_unknown_filter_includes_both_states(client):
    response = client.get("/project_search?drawings=yes_or_unknown")
    assert response.status_code == 200
    assert b"=Formula project" in response.data
    assert b"Library project" in response.data
    assert b"Library remodel" not in response.data
    assert client.get("/project_search?drawings=unknown").status_code == 400


def test_campus_client_is_not_a_searchable_project_field(client):
    response = client.get("/project_search?query=Campus-only")
    assert response.status_code == 200
    assert b"No matching project records were found" in response.data


def test_exact_linked_caan_code_matches_without_partial_code_matches(app):
    with app.app_context():
        from werkzeug.datastructures import MultiDict
        from archives_application.project_tools.project_search import parse_request, search_html

        exact_results, _ = search_html(parse_request(MultiDict([("query", "  70a  ")])) )
        partial_results, _ = search_html(parse_request(MultiDict([("query", "70")])) )
        assert [result.project.id for result in exact_results] == [1]
        assert exact_results[0].ranking_band == "Exact linked CAAN"
        assert exact_results[0].matched_in == "CAAN"
        assert partial_results == []

        mixed_results, _ = search_html(parse_request(MultiDict([("query", "7110 Windows")])) )
        assert [result.project.id for result in mixed_results] == [3]
        assert mixed_results[0].matched_in == "Contract, CAAN"


def test_funding_numbers_match_exact_whitespace_delimited_tokens(app):
    with app.app_context():
        from werkzeug.datastructures import MultiDict
        from archives_application.project_tools.project_search import parse_request, search_html

        exact_results, _ = search_html(parse_request(MultiDict([("query", "wo00423905")])) )
        partial_results, _ = search_html(parse_request(MultiDict([("query", "WO004")])) )
        assert [result.project.id for result in exact_results] == [3]
        assert partial_results == []


def test_html_result_limit_is_applied_after_stable_ranking(app):
    with app.app_context():
        db.session.add_all(
            ProjectModel(
                id=100 + index,
                number=f"9{index:03d}",
                name=f"Open project {index}",
                closed=False,
            )
            for index in range(301)
        )
        db.session.commit()

        from werkzeug.datastructures import MultiDict
        from archives_application.project_tools.project_search import parse_request, search_html

        results, has_more = search_html(parse_request(MultiDict([("status", "open")])) )
        assert len(results) == 300
        assert has_more
        assert results[0].project.id == 1


def test_export_flattens_contracts_and_neutralizes_formula_text(client):
    response = client.get("/project_search/export?query=1000")
    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.data), read_only=True, data_only=False)
    assert workbook.sheetnames == ["Projects and contracts", "Search information"]
    rows = list(workbook["Projects and contracts"].iter_rows(values_only=True))
    assert len(rows) == 3  # Header plus one row for each duplicate project.
    assert rows[0][3] == "Project Information URL"
    assert rows[0][-2:] == ("Ranking band", "database index")
    assert rows[1][1] == "1000"
    assert rows[1][3].startswith("http://")
    assert rows[1][3].endswith("/project_info?project_id=1")
    assert rows[1][2] == "'=Formula project"
    assert rows[1][15] is None  # A project with no contracts still has one blank row.
    assert rows[1][-1] == 1
    assert rows[2][15] == "A-10" and rows[2][-1] == 2
    info = dict(workbook["Search information"].iter_rows(min_row=2, values_only=True))
    assert info["Search results URL"].endswith("/project_search?query=1000")
    assert info["Total matched projects"] == 2
    assert info["Total exported rows"] == 2
